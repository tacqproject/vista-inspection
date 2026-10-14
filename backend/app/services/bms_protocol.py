"""
BMS inspection protocol ("Protokol o běžné/hlavní prohlídce mostu") as .docx —
laid out exactly like the ŘSD BMS protocol the customers submit (reference:
BPM_16-045a_8_7_2026_493d.docx): title page, header data, A. Základní údaje,
C. Stav a závady částí mostu (grouped by bridge part), D, E. Opatření grouped by
urgency, F, H and J. Obrazové přílohy (3 photos per page, caption strip on the
right, linked defect in bold).

The PDF variant is rendered FROM this same .docx (LibreOffice), so both look
identical.
"""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
import unicodedata
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Twips

from app.config import settings

# ─── Catalogues ───────────────────────────────────────────────────────────────

# Parts of a bridge as used by the BMS inspection protocol (code → name).
BRIDGE_PARTS: list[tuple[str, str]] = [
    ("1.1", "Základy"),
    ("1.2", "Mostní podpěry, křídla a čelní zdi"),
    ("1.3", "Úložné prahy, kvádry"),
    ("1.4", "Svahy, kužely, dlažby"),
    ("2.1", "Nosná konstrukce"),
    ("2.2", "Ložiska, klouby"),
    ("2.3", "Mostní závěry"),
    ("3.1", "Vozovka na mostě"),
    ("3.2", "Izolace"),
    ("3.3", "Římsy, obrubníky, zálivky"),
    ("3.4", "Odvodnění"),
    ("3.5", "Přechodové oblasti"),
    ("4.1", "Svodidla/Zábradelní svodidla"),
    ("4.2", "Zábradlí"),
    ("4.3", "Revizní zařízení"),
    ("4.4", "Dopravní značení, osvětlení"),
    ("4.5", "Inženýrské sítě na mostě"),
    ("4.6", "Území pod mostem a přístup. cesty"),
]
PART_GROUPS = {
    "1": "Spodní stavba",
    "2": "Nosná konstrukce mostu (horní stavba)",
    "3": "svršek",
    "4": "Vybavení",
}
# Order of the urgency groups in section E (labels are printed as stored).
URGENCY_ORDER = [
    "periodicky",
    "odstranění nutno ihned",
    "odstranění nutno do 1 roku",
    "odstranění nutno do 2 let",
    "odstranění nutno do 5 let",
    "v rámci běžné údržby",
]
KINDS = {
    "bezna": ("BPM", "BĚŽNÁ PROHLÍDKA"),
    "hlavni": ("HPM", "HLAVNÍ PROHLÍDKA"),
    "prvni_hlavni": ("1.HPM", "PRVNÍ HLAVNÍ PROHLÍDKA"),
    "mimoradna": ("MPM", "MIMOŘÁDNÁ PROHLÍDKA"),
}
DEFECT_LABELS = {
    "crack": "Trhlina", "spalling": "Odpadávání krycí vrstvy", "stain": "Skvrna/výkvět",
    "moss": "Vegetace", "corrosion": "Koroze", "deformation": "Deformace",
    "moisture": "Vlhkost/zatékání", "other": "Jiná závada",
}

# ─── Layout constants (measured from the reference document) ─────────────────

FONT = "Arial"
GREY = RGBColor(0x33, 0x33, 0x33)
SZ_BODY = Pt(10.5)        # w:sz 21
SZ_NOTE = Pt(9.5)         # w:sz 19
SZ_CAPTION = Pt(9)        # w:sz 18
PAGE_W, PAGE_H = Twips(11900), Twips(16840)
PHOTO_W, PHOTO_H = Emu(3591673), Emu(2694929)          # 9.98 × 7.49 cm, 4:3
PHOTO_ROW_H = Twips(4422)                               # 7.80 cm → 3 per page


# ─── Low-level helpers ────────────────────────────────────────────────────────

def _set_font(run, size=SZ_BODY, bold=False, color=None, italic=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _para(container, text="", size=SZ_BODY, bold=False, color=None, after=0, before=0,
          left=0, align=None, line=1.08, keep_next=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after); pf.space_before = Pt(before)
    pf.line_spacing = line
    if left:
        pf.left_indent = Twips(left)
    if align is not None:
        p.alignment = align
    if keep_next:
        pf.keep_with_next = True
    if text:
        _set_font(p.add_run(text), size, bold, color)
    return p


def _heading(doc, text, level=2, before=12):
    """Nadpis2 = 14 pt bold (A.–J. sections); Nadpis3 = 11 pt bold (part groups)."""
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(5.55 if level == 2 else 6)
    p.paragraph_format.keep_with_next = True
    _set_font(p.add_run(text), Pt(14) if level == 2 else Pt(11), True, RGBColor(0, 0, 0))
    return p


def _field(run_parent, instr: str):
    """Append a Word field (PAGE / NUMPAGES) to a paragraph."""
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        r = run_parent.add_run()
        _set_font(r, Pt(11))
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {text} "
            r._element.append(it)
        elif kind == "text":
            r.text = text
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind)
            r._element.append(fc)


def _p_border(p, side="bottom", color="000000", sz=6, space=1):
    ppr = p._element.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr"); ppr.append(bdr)
    b = OxmlElement(f"w:{side}")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", str(space)), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)


def _p_shade(p, fill="EEEEEE"):
    ppr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _table_borders(table, spec: dict[str, tuple[str, int] | None]):
    """spec: side → (color, size in 1/8 pt) or None for no border."""
    tblpr = table._element.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        v = spec.get(side)
        if v is None:
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v[1]))
            e.set(qn("w:space"), "0"); e.set(qn("w:color"), v[0])
        b.append(e)
    tblpr.append(b)


def _col_widths(table, widths_twips: list[int]):
    table.autofit = False
    tblpr = table._element.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(sum(widths_twips))); w.set(qn("w:type"), "dxa")
    old = tblpr.find(qn("w:tblW"))
    if old is not None:
        tblpr.remove(old)
    tblpr.append(w)
    grid = table._element.tblGrid
    for gc, wt in zip(grid.findall(qn("w:gridCol")), widths_twips):
        gc.set(qn("w:w"), str(wt))
    for row in table.rows:
        for cell, wt in zip(row.cells, widths_twips):
            cell.width = Twips(wt)


def _cell_margins(table, left=0, right=0, top=0, bottom=0):
    tblpr = table._element.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblpr.append(m)


def _cell_text(cell, text, size=SZ_BODY, bold=False, color=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    lines = (text or "").split("\n")
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
        _set_font(p.add_run(line), size, bold, color)


def _ascii(s: str) -> str:
    return unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode()


def _fmt_date(d) -> str:
    if not d:
        return ""
    if isinstance(d, str):
        try:
            d = date.fromisoformat(d[:10])
        except ValueError:
            return d
    return f"{d.day}.{d.month}.{d.year}"


def _v(x) -> str:
    return "" if x is None else str(x)


def _part_label(a: dict) -> str:
    bp = (a.get("bridge_part") or "").strip()
    if bp:
        return bp
    return DEFECT_LABELS.get(a.get("defect_type", ""), a.get("defect_type") or "Nezařazeno")


def _part_group(a: dict) -> str:
    bp = (a.get("bridge_part") or "").strip()
    head = bp.split(" ", 1)[0].split(".", 1)[0] if bp else ""
    return head if head in PART_GROUPS else "9"


def _part_sort(a: dict):
    bp = (a.get("bridge_part") or "").strip().split(" ", 1)[0]
    try:
        return tuple(int(x) for x in bp.split(".") if x)
    except ValueError:
        return (99,)


# ─── Photos ───────────────────────────────────────────────────────────────────

_IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")


def appendix_dir(inspection_id: str) -> Path:
    return settings.upload_path / "inspections" / inspection_id / "report_photos"


def _photo_box(path: Path) -> io.BytesIO | None:
    """EXIF-upright photo fitted into the protocol's 4:3 frame (white letterbox for
    other aspect ratios — nothing of the evidence is cropped), downsized so a
    60-photo protocol stays a few tens of MB."""
    try:
        from PIL import Image, ImageOps
        im = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    except Exception:
        return None
    W, H = 1600, 1200
    im.thumbnail((W, H), Image.LANCZOS)
    if abs(im.width / im.height - 4 / 3) < 0.02:
        canvas = im.resize((W, H), Image.LANCZOS)
    else:
        canvas = Image.new("RGB", (W, H), "white")
        canvas.paste(im, ((W - im.width) // 2, (H - im.height) // 2))
    bio = io.BytesIO()
    canvas.save(bio, "JPEG", quality=85)
    bio.seek(0)
    return bio


def _outline_picture(run):
    """0.25 mm black frame around the photo (as in the reference)."""
    pics = run._element.findall(".//" + qn("pic:spPr"))
    for sppr in pics:
        ln = OxmlElement("a:ln"); ln.set("w", "9144")
        fill = OxmlElement("a:solidFill"); clr = OxmlElement("a:srgbClr"); clr.set("val", "000000")
        fill.append(clr); ln.append(fill); sppr.append(ln)


def _natural(name: str):
    import re
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", name)]


def _caption(name: str) -> str:
    """'01 Pohled ve směru staničení.jpg' → 'Pohled ve směru staničení.jpg': a leading
    number only orders the photos; the protocol numbers them itself."""
    import re
    return re.sub(r"^\d+\s*[._\-)]?\s*", "", name) or name


def collect_photos(inspection_id: str, defects: list[dict]) -> list[dict]:
    """General appendix photos first (sorted by filename), then each defect's
    detail photos in section-C order, each linked to its defect."""
    items: list[dict] = []
    gdir = appendix_dir(inspection_id)
    if gdir.is_dir():
        for f in sorted(gdir.iterdir(), key=lambda f: _natural(f.name)):
            if f.is_file() and f.suffix.lower() in _IMG_EXT:
                items.append({"path": f, "caption": _caption(f.name), "defect": None})
    base = settings.upload_path / "inspections" / inspection_id / "annotation_images"
    for a in defects:
        d = base / a["id"]
        if not d.is_dir():
            continue
        for f in sorted(d.iterdir(), key=lambda f: _natural(f.name)):
            if f.is_file() and f.suffix.lower() in _IMG_EXT:
                items.append({"path": f, "caption": _caption(f.name), "defect": a})
    return items


# ─── Document ─────────────────────────────────────────────────────────────────

def _setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = SZ_BODY
    st.element.get_or_add_rPr()
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.08
    for lvl in (1, 2, 3):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = FONT
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.font.size = {1: Pt(23), 2: Pt(14), 3: Pt(11)}[lvl]
        rpr = h.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]


def _setup_section(sec, top, right, bottom, left):
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.top_margin, sec.right_margin = Twips(top), Twips(right)
    sec.bottom_margin, sec.left_margin = Twips(bottom), Twips(left)
    sec.header_distance, sec.footer_distance = Twips(643), Twips(624)


def build_protocol(structure: dict, inspection: dict, annotations: list[dict],
                   extra_views: dict | None = None) -> bytes:
    """Build the protocol .docx. `extra_views` optionally maps annotation id →
    data-URL of a 3D close-up, appended to the photo appendix."""
    doc = Document()
    _setup_styles(doc)
    sec = doc.sections[0]
    _setup_section(sec, 1421, 1440, 1440, 1440)

    kind_code, kind_title = KINDS.get(inspection.get("inspection_kind") or "bezna", KINDS["bezna"])
    ev = structure.get("evidence_number") or ""
    name = structure.get("name") or ""
    is_bridge = (structure.get("type") or "").lower() in ("bridge", "most", "lavka", "propustek") or bool(ev)
    obj_word = "Most" if is_bridge else "Objekt"
    insp_date = _fmt_date(inspection.get("inspection_date"))
    inspector = inspection.get("inspector_name") or ""

    # ── header / footer (shared by all sections) ──────────────────────────────
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = ", ".join(x for x in (insp_date, _ascii(inspector)) if x)
    _set_font(hp.add_run(f"{kind_code} {ev or _ascii(name)}" + (f" ({meta})" if meta else "")), Pt(11))
    _p_border(hp, "bottom", "000000", 6, 1)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(fp.add_run("Strana "), Pt(11)); _field(fp, "PAGE   \\* MERGEFORMAT")
    _set_font(fp.add_run(" z "), Pt(11)); _field(fp, "NUMPAGES   \\* MERGEFORMAT")

    # ── title page ─────────────────────────────────────────────────────────────
    _para(doc, f"{obj_word} {ev}".strip() if ev else name, Pt(27.5), True,
          align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    if ev:
        _para(doc, name, Pt(15.5), after=20.5, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    t = doc.add_paragraph(style="Heading 1")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(t.add_run(kind_title), Pt(23), True, RGBColor(0, 0, 0))
    t.add_run().add_break(WD_BREAK.PAGE)

    # ── identification ────────────────────────────────────────────────────────
    obj = f"Objekt: {obj_word} ev.č. {ev} ({name})" if ev else f"Objekt: {name}"
    _set_font(_para(doc).add_run(obj), Pt(11), True)
    if structure.get("district"):
        _para(doc, f"Okres: {structure['district']}", after=5.3)
    _para(doc, f"Prohlídku provedl:  {inspector}")
    if inspection.get("inspector_company"):
        _para(doc, f" {inspection['inspector_company']}", left=1570)
    _para(doc, f"Datum provedení prohlídky: {insp_date}")
    _para(doc, "Poznámka:")
    for line in (inspection.get("protocol_note") or "").splitlines():
        _para(doc, line, SZ_NOTE, color=GREY, left=553, line=1.1)
    _para(doc, "Počasí v době provádění prohlídky:")
    _para(doc, _v(inspection.get("weather")), SZ_NOTE, color=GREY, left=553, line=1.1)
    _para(doc, "Způsob zpřístupnění:")
    _para(doc, _v(inspection.get("access_method")), SZ_NOTE, color=GREY, left=553, line=1.1)
    p = _para(doc)
    at = _v(inspection.get("air_temperature"))
    at = f"{at}°C" if at and not at.endswith("C") else at
    _set_font(p.add_run(f"Teplota vzduchu: {at}\tTeplota NK: {_v(inspection.get('structure_temperature'))}"))
    p.paragraph_format.tab_stops.add_tab_stop(Twips(4800))

    # ── A. Základní údaje ─────────────────────────────────────────────────────
    _heading(doc, "A. ZÁKLADNÍ ÚDAJE", before=24)
    p = _para(doc)
    p.paragraph_format.tab_stops.add_tab_stop(Twips(2600))
    p.paragraph_format.tab_stops.add_tab_stop(Twips(5600))
    _set_font(p.add_run(f"Číslo komunikace: {_v(structure.get('road_number'))}\t"
                        f"Staničení km: {_v(structure.get('chainage_km'))}"
                        f"{'km' if structure.get('chainage_km') and 'km' not in str(structure.get('chainage_km')) else ''}\t"
                        f"Ev.č.mostu: {ev}"))
    p = _para(doc)
    _set_font(p.add_run("Název objektu: "))
    _set_font(p.add_run(name), bold=True)
    p = _para(doc)
    for pos in (1400, 2800, 4200):
        p.paragraph_format.tab_stops.add_tab_stop(Twips(pos))
    _set_font(p.add_run(f"Vn= {_v(structure.get('load_vn'))}\tVr= {_v(structure.get('load_vr'))}\t"
                        f"Ve= {_v(structure.get('load_ve'))}\tMax.nápr.tlak = {_v(structure.get('max_axle_load'))}"))
    p = _para(doc)
    p.paragraph_format.tab_stops.add_tab_stop(Twips(1400))
    p.paragraph_format.tab_stops.add_tab_stop(Twips(4200))
    _set_font(p.add_run(f"Stav mostu:\tspodní stavba: {_v(inspection.get('state_substructure'))}\t"
                        f"nosná konstrukce: {_v(inspection.get('state_superstructure'))}"))
    _para(doc, f"Použitelnost: {_v(inspection.get('usability'))}")
    _para(doc, f"Staničení ve směru: {_v(structure.get('chainage_direction'))}")

    # ── C. Stav a závady částí mostu ──────────────────────────────────────────
    defects = sorted(annotations, key=lambda a: (_part_group(a), _part_sort(a), str(a.get("created_at") or "")))
    _heading(doc, "C. STAV A ZÁVADY ČÁSTÍ MOSTU", before=24)
    groups: dict[str, list[dict]] = {}
    for a in defects:
        groups.setdefault(_part_group(a), []).append(a)
    for g in sorted(groups):
        _heading(doc, f"{g}. {PART_GROUPS.get(g, 'Ostatní')}", level=3, before=6)
        tbl = doc.add_table(rows=0, cols=3)
        _table_borders(tbl, {})
        _cell_margins(tbl, left=0, right=60, top=0, bottom=90)
        for i, a in enumerate(groups[g], 1):
            row = tbl.add_row()
            row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _cell_text(row.cells[0], f"[{g}.{i}]")
            _cell_text(row.cells[1], _part_label(a))
            desc = (a.get("description") or "").strip() or DEFECT_LABELS.get(a.get("defect_type", ""), "")
            _cell_text(row.cells[2], desc, color=GREY)
        _col_widths(tbl, [1061, 2700, 5265])

    # ── D ─────────────────────────────────────────────────────────────────────
    _heading(doc, "D. HODNOCENÍ PÉČE O MOST, VÝKONU BĚŽNÝCH PROHLÍDEK, KVALITY ÚDRŽBOVÝCH "
                  "PRACÍ A PROVÁDĚNÝCH OPRAV, ZÁVADY MOSTNÍ EVIDENCE", before=24)
    for line in (inspection.get("care_assessment") or "").splitlines() or [""]:
        _para(doc, line, SZ_NOTE, color=GREY, line=1.1)

    # ── E. Opatření by urgency ────────────────────────────────────────────────
    _heading(doc, "E. OPATŘENÍ NA ZKVALITNĚNÍ SPRÁVY MOSTU, NÁVRH NA ODSTRANĚNÍ "
                  "ZJIŠTĚNÝCH ZÁVAD", before=24)
    measures = [a for a in defects if (a.get("recommended_action") or "").strip()]

    def urg_key(u: str):
        u = (u or "").strip()
        low = u.lower()
        for i, k in enumerate(URGENCY_ORDER):
            if k in low:
                return (i, u)
        return (len(URGENCY_ORDER), u)

    by_urg: dict[str, list[dict]] = {}
    for a in measures:
        by_urg.setdefault((a.get("action_urgency") or "bez určení termínu").strip(), []).append(a)
    n = 0
    for u in sorted(by_urg, key=urg_key):
        _para(doc, u, SZ_BODY, bold=True, before=6, after=3, keep_next=True)
        tbl = doc.add_table(rows=0, cols=3)
        _table_borders(tbl, {})
        _cell_margins(tbl, left=0, right=60, top=0, bottom=90)
        for a in by_urg[u]:
            n += 1
            row = tbl.add_row()
            _cell_text(row.cells[0], f"[{n}]")
            _cell_text(row.cells[1], _part_label(a))
            _cell_text(row.cells[2], a.get("recommended_action") or "")
        _col_widths(tbl, [700, 2800, 5526])
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tblpr = tbl._element.tblPr
        ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "220"); ind.set(qn("w:type"), "dxa")
        tblpr.append(ind)

    # ── F ─────────────────────────────────────────────────────────────────────
    _heading(doc, "F. ZÁZNAM O PROJEDNÁNÍ OPATŘENÍ SE SPRÁVCEM MOSTU, STANOVENÍ DRUHU ÚDRŽBY "
                  "A OPRAV, STANOVENÍ ZPŮSOBU A TERMÍNU ODSTRANĚNÍ ZÁVAD, PŘÍPADNÉ NAŘÍZENÍ "
                  "ZATĚŽOVACÍ ZKOUŠKY, STANOVENÍ PŘEDBĚŽNÉ CENY PRACÍ", before=24)
    _para(doc, f"Datum projednání: {_fmt_date(inspection.get('negotiation_date'))}")
    _para(doc, f"Číslo jednací: {_v(inspection.get('reference_number'))}", after=5.35)
    _para(doc, "Poznámka:")
    for line in (inspection.get("negotiation_note") or "").splitlines() or [""]:
        _para(doc, line, SZ_NOTE, color=GREY, left=553, line=1.1)

    # ── H ─────────────────────────────────────────────────────────────────────
    _heading(doc, "H. STANOVISKO NADŘÍZENÉHO ORGÁNU K PŘÍPADNÝM POŽADAVKŮM SPRÁVCE MOSTU",
             before=24)
    for line in (inspection.get("authority_opinion") or "").splitlines():
        _para(doc, line, SZ_NOTE, color=GREY, line=1.1)

    # ── J. Obrazové přílohy (own section: the reference uses wider margins) ───
    photos = collect_photos(inspection["id"], defects)
    views = extra_views or {}
    for a in defects:
        data = views.get(a.get("id"))
        if data:
            photos.append({"data_url": data, "caption": f"3D pohled – {_part_label(a)}", "defect": a})

    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    _setup_section(sec2, 1435, 892, 1547, 1535)
    _heading(doc, "J. OBRAZOVÉ PŘÍLOHY", before=0).paragraph_format.space_after = Pt(0.35)
    if photos:
        tbl = doc.add_table(rows=0, cols=2)
        _table_borders(tbl, {"insideH": ("CCCCCC", 2), "top": ("CCCCCC", 2)})
        _cell_margins(tbl, left=0, right=0, top=0, bottom=0)
        for i, ph in enumerate(photos, 1):
            row = tbl.add_row()
            row.height = PHOTO_ROW_H
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            trpr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trpr.append(cs)
            c0, c1 = row.cells
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_before = Pt(1.5)
            if "path" in ph:
                img = _photo_box(ph["path"])
            else:
                import base64
                raw = ph["data_url"].split(",", 1)[-1]
                img = io.BytesIO(base64.b64decode(raw))
            if img is not None:
                run = p0.add_run()
                run.add_picture(img, width=PHOTO_W, height=PHOTO_H)
                _outline_picture(run)
            # caption strip + linked defect
            cp = c1.paragraphs[0]
            cp.paragraph_format.left_indent = Twips(150)
            cp.paragraph_format.space_before = Pt(6)
            cp.paragraph_format.space_after = Pt(3)
            _p_shade(cp, "EEEEEE")
            _set_font(cp.add_run(f"{i:02d}.{ph['caption']}"), SZ_CAPTION)
            a = ph.get("defect")
            if a:
                q = c1.add_paragraph(); q.paragraph_format.left_indent = Twips(150)
                q.paragraph_format.space_after = Pt(0)
                _set_font(q.add_run(_part_label(a)), SZ_CAPTION, True)
                desc = (a.get("description") or "").strip()
                for line in desc.splitlines():
                    q = c1.add_paragraph(); q.paragraph_format.left_indent = Twips(150)
                    q.paragraph_format.space_after = Pt(0)
                    q.paragraph_format.line_spacing = 1.1
                    _set_font(q.add_run(line), SZ_CAPTION, color=GREY)
        _col_widths(tbl, [5688, 3785])
    else:
        _para(doc, "Bez obrazových příloh.", SZ_NOTE, color=GREY, before=6)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_to_pdf(docx_bytes: bytes) -> bytes:
    """Render the .docx to PDF with LibreOffice (same layout as the .docx)."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice (soffice) není nainstalován — PDF nelze vytvořit")
    with tempfile.TemporaryDirectory(prefix="vista_pdf_") as tmp:
        src = Path(tmp) / "protokol.docx"
        src.write_bytes(docx_bytes)
        profile = Path(tmp) / "lo_profile"          # private profile → parallel-safe
        subprocess.run(
            [exe, f"-env:UserInstallation=file://{profile}", "--headless",
             "--convert-to", "pdf", "--outdir", tmp, str(src)],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300,
        )
        return (Path(tmp) / "protokol.pdf").read_bytes()


def protocol_filename(structure: dict, inspection: dict, ext: str) -> str:
    kind_code, _ = KINDS.get(inspection.get("inspection_kind") or "bezna", KINDS["bezna"])
    ev = structure.get("evidence_number") or _ascii(structure.get("name") or "objekt")
    d = inspection.get("inspection_date")
    ds = f"{d.day}_{d.month}_{d.year}" if isinstance(d, (date, datetime)) else _ascii(str(d or ""))
    safe = "".join(c if c.isalnum() or c in "-_." else "_" for c in _ascii(f"{kind_code}_{ev}_{ds}"))
    return f"{safe}.{ext}"
